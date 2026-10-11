# AutoQgen_App/services.py

import io
import docx
import requests
import pandas as pd
import cloudinary.uploader
from lxml import etree
import latex2mathml.converter

NAMESPACES = {
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'v': 'urn:schemas-microsoft-com:vml'
}

# MathML ko Microsoft Word ke native OMML format me convert karne ka stylesheet
MML2OMML_XSL = etree.XML('''
<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform" 
xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" 
xmlns:mml="http://www.w3.org/1998/Math/MathML">
  <xsl:template match="mml:math">
    <m:oMath>
      <xsl:apply-templates select="*"/>
    </m:oMath>
  </xsl:template>
  <xsl:template match="mml:mfrac">
    <m:f>
      <m:num><m:e><xsl:apply-templates select="*[1]"/></m:e></m:num>
      <m:den><m:e><xsl:apply-templates select="*[2]"/></m:e></m:den>
    </m:f>
  </xsl:template>
  <xsl:template match="mml:msqrt">
    <m:rad>
      <m:radPr><m:degHide m:val="1"/></m:radPr>
      <m:deg/>
      <m:e><xsl:apply-templates select="*"/></m:e>
    </m:rad>
  </xsl:template>
  <xsl:template match="mml:msup">
    <m:sSup>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sup><xsl:apply-templates select="*[2]"/></m:sup>
    </m:sSup>
  </xsl:template>
  <xsl:template match="mml:msub">
    <m:sSub>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sub><xsl:apply-templates select="*[2]"/></m:sub>
    </m:sSub>
  </xsl:template>
  <xsl:template match="mml:msubsup">
    <m:sSubSup>
      <m:e><xsl:apply-templates select="*[1]"/></m:e>
      <m:sub><xsl:apply-templates select="*[2]"/></m:sub>
      <m:sup><xsl:apply-templates select="*[3]"/></m:sup>
    </m:sSubSup>
  </xsl:template>
  <xsl:template match="mml:mi | mml:mn | mml:mo">
    <m:r>
      <m:t><xsl:value-of select="."/></m:t>
    </m:r>
  </xsl:template>
</xsl:stylesheet>
''')

def to_pure_lxml(element):
    """BaseOxmlElement ko pure lxml element me convert karta hai"""
    return etree.fromstring(element.xml.encode('utf-8'))

def extract_math(element):
    """OMML math elements ko readable text/latex me nikalta hai."""
    root = to_pure_lxml(element)
    math_elements = root.xpath('.//m:oMathPara | .//m:oMath[not(ancestor::m:oMathPara)]', namespaces=NAMESPACES)
    eqs = []
    for m in math_elements:
        math_texts = m.xpath('.//m:t', namespaces=NAMESPACES)
        eq_text = "".join([t.text for t in math_texts if t.text])
        if eq_text.strip():
            eqs.append(f"${eq_text.strip()}$")
    return " | ".join(eqs)

def extract_and_upload_images(element, doc_part):
    """Drawing blips nikal kar Cloudinary par upload karta hai."""
    root = to_pure_lxml(element)
    image_urls = []

    # 1. Standard DrawingML Blips
    blips = root.xpath('.//a:blip', namespaces=NAMESPACES)
    for blip in blips:
        r_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if r_id and r_id in doc_part.rels:
            rel = doc_part.rels[r_id]
            if "image" in rel.target_ref:
                image_bytes = rel.target_part.blob
                upload_res = cloudinary.uploader.upload(
                    io.BytesIO(image_bytes),
                    folder="exam_diagrams",
                    resource_type="image"
                )
                image_urls.append(upload_res.get("secure_url"))

    # 2. Legacy VML Drawings
    vml_images = root.xpath('.//v:imagedata', namespaces=NAMESPACES)
    for vml in vml_images:
        r_id = vml.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        if r_id and r_id in doc_part.rels:
            rel = doc_part.rels[r_id]
            if "image" in rel.target_ref:
                image_bytes = rel.target_part.blob
                upload_res = cloudinary.uploader.upload(
                    io.BytesIO(image_bytes),
                    folder="exam_diagrams",
                    resource_type="image"
                )
                image_urls.append(upload_res.get("secure_url"))

    return " | ".join(image_urls)

def process_docx_file(uploaded_file):
    """DOCX file se paragraphs aur tables parse karke DataFrame banata hai."""
    doc = docx.Document(io.BytesIO(uploaded_file.read()))
    records = []
    idx = 1

    # Paragraphs scan
    for p in doc.paragraphs:
        text = p.text.strip()
        p_elm = p._element
        eqs = extract_math(p_elm)
        imgs = extract_and_upload_images(p_elm, doc.part)

        if text or eqs or imgs:
            records.append({
                "item_id": idx,
                "type": "paragraph",
                "content": text,
                "equations": eqs,
                "diagram_urls": imgs
            })
            idx += 1

    # Tables scan
    for t_i, table in enumerate(doc.tables):
        for r_i, row in enumerate(table.rows):
            for c_i, cell in enumerate(row.cells):
                c_text = cell.text.strip()
                cell_elm = cell._element
                eqs = extract_math(cell_elm)
                imgs = extract_and_upload_images(cell_elm, doc.part)
                if c_text or eqs or imgs:
                    records.append({
                        "item_id": idx,
                        "type": f"table_{t_i}_r{r_i}_c{c_i}",
                        "content": c_text,
                        "equations": eqs,
                        "diagram_urls": imgs
                    })
                    idx += 1

    return pd.DataFrame(records)

def upload_dataframe_as_excel(df, public_id):
    """DataFrame ko Excel memory me convert karke Cloudinary par overwrite karta hai."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
    output.seek(0)

    upload_result = cloudinary.uploader.upload(
        output,
        resource_type="raw",
        public_id=f"{public_id}.xlsx",
        overwrite=True,
        invalidate=True
    )
    return upload_result['secure_url'], upload_result['public_id']

def add_math_to_docx_paragraph(paragraph, latex_str):
    """
    LaTeX expression ko Word ke native Equation object (<m:oMath>)
    me convert karke paragraph me embed karta hai.
    """
    clean_latex = latex_str.replace('$', '').strip()
    if not clean_latex:
        return

    try:
        # LaTeX -> MathML
        mathml_str = latex2mathml.converter.convert(clean_latex)
        
        # MathML -> OMML
        mathml_tree = etree.fromstring(mathml_str.encode('utf-8'))
        transform = etree.XSLT(MML2OMML_XSL)
        omml_tree = transform(mathml_tree)

        # Word Document Paragraph me append
        paragraph._p.append(omml_tree.getroot())
    except Exception:
        # Agar koi symbol parse na ho paye toh fallback cleanly Cambria Math me text banayega
        r = paragraph.add_run(f" {clean_latex} ")
        r.italic = True
        r.font.name = "Cambria Math"

def generate_exam_docx(doc_title, records):
    """
    Records (JSON/Dict) ko standard Exam Paper .docx me convert karta hai
    jisme Equations live MS Word math elements banengi aur images center aligned embed hongi.
    """
    out_doc = docx.Document()

    # Exam Header Template
    header = out_doc.add_heading(doc_title, level=1)
    header.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    
    meta_p = out_doc.add_paragraph()
    meta_p.add_run("Time: 3 Hours\t\t\t\tFull Marks: 100").bold = True
    out_doc.add_paragraph("―" * 52).alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

    # Questions Render
    for row in records:
        content = str(row.get("content", "")).strip()
        eqs = str(row.get("equations", "")).strip()
        img_urls = str(row.get("diagram_urls", "")).strip()

        if not (content or eqs or img_urls):
            continue

        p = out_doc.add_paragraph()
        if content:
            p.add_run(content + "  ")

        # Native Word Math Equations
        if eqs:
            for single_eq in eqs.split('|'):
                add_math_to_docx_paragraph(p, single_eq.strip())

        # Cloudinary images embed
        if img_urls:
            for url in img_urls.split('|'):
                url = url.strip()
                if url.startswith("http"):
                    try:
                        resp = requests.get(url, timeout=10)
                        if resp.status_code == 200:
                            img_stream = io.BytesIO(resp.content)
                            img_p = out_doc.add_paragraph()
                            img_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                            img_p.add_run().add_picture(img_stream, width=docx.shared.Inches(3.2))
                    except Exception as e:
                        print(f"Error fetching image {url}: {e}")

    output_stream = io.BytesIO()
    out_doc.save(output_stream)
    output_stream.seek(0)
    return output_stream